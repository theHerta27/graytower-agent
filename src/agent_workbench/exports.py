"""Derive readable exports from the exact frozen candidate; never use a model."""
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def bundle_for_task(workbench, task):
    from .artifact_views import frozen_candidate
    candidate = frozen_candidate(workbench.repo, task['candidate_id'])
    bundle = workbench.repo.read_artifact(candidate['manifest']['candidate']['sha256'])
    return candidate, bundle


def workbook_for_task(workbench, task):
    candidate, bundle = bundle_for_task(workbench, task)
    book = Workbook()
    skills = book.active
    skills.title = '技能配置'
    skills.append(['技能 ID', '字段', '值'])
    for name, values in bundle.get('skills', {}).items():
        if isinstance(values, dict):
            for field, value in values.items():
                skills.append([name, field, value if isinstance(value, (str, int, float, bool)) else json.dumps(value, ensure_ascii=False)])
        else:
            skills.append([name, '值', str(values)])
    diff = book.create_sheet('变更对照')
    diff.append(['字段路径', '原值', '新值', '需求', '校验状态'])
    for item in candidate['diff']:
        diff.append([item['path'], item['before'], item['after'], task['request'],
                     '预检通过' if candidate['validation'].get('passed') else '检查结果见候选'])
    version = book.create_sheet('版本说明')
    for row in [('任务号', task['task_id']), ('需求', task['request']), ('候选 Hash', task['candidate_id']),
                ('基线 Hash', candidate['manifest']['baseline']['sha256']), ('任务创建时间', task['created_at']),
                ('来源', task['origin']), ('验证状态', task['status']), ('应用状态', '未应用到游戏'),
                ('权威数据', '冻结 JSON；Excel 为只读派生视图')]:
        version.append(row)
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Formula-like user text must remain literal text.
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='DDEEE9')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(65, max(18, max(len(str(c.value or '')) for c in column)+2))
    output = BytesIO()
    book.save(output)
    return output.getvalue()
