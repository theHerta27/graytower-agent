"""Task/UI contracts with scripted model decisions; no provider calls."""
import io
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from starlette.testclient import TestClient

from agent_workbench.domain import CandidateRepository
from agent_workbench.workbench import create_app


class Model:
    calls = 0
    def remaining_calls(self):
        return 4 - self.calls
    def complete(self, messages, schema):
        self.calls += 1
        if schema.__name__ == 'ConfigDecision':
            return {'status': 'candidate', 'message': '脚本化工程测试输出',
                    'changes': [{'path': 'skills.breaker_skill.multiplier', 'value': 1.6}]}
        return {'action': 'finish', 'reason': '工程测试不执行模拟', 'seeds': [], 'policies': [], 'artifact_hash': ''}


class TaskTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.repo = CandidateRepository(self.root)

    def client(self, model=None):
        args = {} if model is None else {'call_cap': 4, 'audit_path': self.root / 'fixture.jsonl',
                                          'model_factory': lambda **kw: model}
        return TestClient(create_app(self.root, **args), base_url='http://127.0.0.1:8765')

    def body(self, **kw):
        return {'task_id': 'task1', 'kind': 'config', 'request': '将倍率改为1.6，其余保持不变', **kw}

    def wait(self, client, field, expected):
        for _ in range(150):
            task = client.get('/api/tasks/task1').json()
            if task.get(field) == expected or (expected == 'present' and task.get(field)):
                return task
            time.sleep(.03)
        self.fail(str(task))

    def test_disabled_model_records_request_without_spending_and_rejects_changed_retry(self):
        with self.client() as c:
            r = c.post('/api/tasks', json=self.body())
            self.assertEqual(r.status_code, 201)
            self.assertEqual(c.post('/api/tasks', json=self.body()).json()['task_id'], 'task1')
            self.assertEqual(c.post('/api/tasks', json=self.body(request='不同需求')).status_code, 409)
            self.assertEqual(c.post('/api/tasks/task1/generate', json={}).status_code, 403)
            self.assertIsNone(c.get('/api/tasks/task1').json()['candidate_id'])

    def test_graph_candidate_export_test_approval_and_qa_are_one_task(self):
        model = Model()
        with self.client(model) as c:
            self.assertEqual(c.post('/api/tasks', json=self.body()).status_code, 201)
            self.assertEqual(c.post('/api/tasks/task1/generate', json={}).status_code, 202)
            task = self.wait(c, 'candidate_id', 'present')
            manifest = task['candidate_id']
            self.assertEqual(c.post('/api/tasks/task1/generate', json={}).status_code, 202)
            self.assertEqual(model.calls, 1)
            raw = c.get('/api/tasks/task1/export/json')
            self.assertEqual(raw.json()['skills']['breaker_skill']['multiplier'], 1.6)
            book = load_workbook(io.BytesIO(c.get('/api/tasks/task1/export/xlsx').content))
            self.assertEqual(book.sheetnames, ['技能配置', '变更对照', '版本说明'])
            self.assertEqual(book['变更对照']['C2'].value, 1.6)
            self.assertEqual(self.repo.get(manifest)['manifest_hash'], manifest)
            self.assertEqual(c.post('/api/tasks/task1/verify', json={'goal': '检查', 'max_battles': 4}).status_code, 409)
            self.assertEqual(c.post('/api/tasks/task1/approve-test', json={}).status_code, 200)
            self.assertEqual(c.post('/api/tasks/task1/verify', json={'goal': '检查', 'max_battles': 4}).status_code, 202)
            task = self.wait(c, 'status', 'verification_completed')
            self.assertEqual(task['report']['request']['manifest_hash'], manifest)
            self.assertEqual(task['report']['request']['report']['verdict'], 'INCONCLUSIVE')
            self.assertTrue(any(e['kind'] == 'test_approved' for e in task['events']))
            calls = model.calls
            c.post('/api/tasks', json=self.body(task_id='history1', candidate_id=manifest))
            historic = c.post('/api/tasks/history1/history', json={'request_id': 'task1_qa'})
            self.assertEqual(historic.status_code, 200)
            self.assertEqual(historic.json()['status'], 'history_view')
            self.assertEqual(model.calls, calls)
            self.assertEqual(c.post('/api/tasks/history1/verify', json={'goal': '检查', 'max_battles': 4}).status_code, 409)
            self.assertEqual(c.post('/api/tasks/history1/history', json={'request_id': 'missing'}).status_code, 404)

    def test_export_request_formula_is_text_and_existing_candidate_is_marked_historical(self):
        candidate = self.repo.create('合法倍率', [{'path': 'skills.breaker_skill.multiplier', 'value': 1.6}])
        with self.client() as c:
            r = c.post('/api/tasks', json=self.body(request='=HYPERLINK("https://example.org")',
                       candidate_id=candidate['manifest_hash']))
            self.assertEqual(r.status_code, 201)
            self.assertEqual(r.json()['origin'], 'existing_candidate')
            book = load_workbook(io.BytesIO(c.get('/api/tasks/task1/export/xlsx').content))
            cells = list(book['版本说明'].iter_rows())
            request_cell = next(row[1] for row in cells if row[0].value == '需求')
            self.assertEqual(request_cell.data_type, 's')
            self.assertEqual(request_cell.value, '=HYPERLINK("https://example.org")')
            self.assertEqual(c.post('/api/tasks/task1/verify', json={'goal': '检查', 'max_battles': 4}).status_code, 409)

    def test_foreign_candidate_and_unknown_body_rejected(self):
        with self.client() as c:
            self.assertEqual(c.post('/api/tasks', json=self.body(candidate_id='f'*64)).status_code, 422)
            self.assertEqual(c.post('/api/tasks', json=self.body(automatic_publish=True)).status_code, 422)
            self.assertEqual(c.get('/api/tasks/task1').status_code, 404)

    def test_rejected_qa_admission_keeps_candidate_ready(self):
        candidate = self.repo.create('合法倍率', [{'path': 'skills.breaker_skill.multiplier', 'value': 1.6}])
        with self.client() as c:
            c.post('/api/tasks', json=self.body(candidate_id=candidate['manifest_hash']))
            c.post('/api/tasks/task1/approve-test', json={})
            self.assertEqual(c.post('/api/tasks/task1/verify', json={'goal': '检查', 'max_battles': 4}).status_code, 403)
            self.assertEqual(c.get('/api/tasks/task1').json()['status'], 'ready_for_verification')

    def test_task_page_can_be_opened_directly(self):
        with self.client() as c:
            self.assertEqual(c.get('/tasks').status_code, 200)

    def test_environment_drift_allows_historical_export_but_blocks_new_qa(self):
        candidate = self.repo.create('合法倍率', [{'path': 'skills.breaker_skill.multiplier', 'value': 1.6}])
        with self.client() as c, patch.object(CandidateRepository, 'get', side_effect=ValueError('environment fingerprint/content drift')):
            r = c.post('/api/tasks', json=self.body(candidate_id=candidate['manifest_hash']))
            self.assertEqual(r.status_code, 201)
            self.assertFalse(r.json()['candidate']['execution_compatible'])
            self.assertEqual(c.get('/api/tasks/task1/export/json').json()['skills']['breaker_skill']['multiplier'], 1.6)
            self.assertEqual(c.post('/api/tasks/task1/approve-test', json={}).status_code, 409)
